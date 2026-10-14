"""Readers that turn an export file into raw rows with their source locators."""

import csv
import datetime
import io
import json

from openpyxl import load_workbook

from ignis_wo.contracts import Locator, RawRow


def cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value).strip()


def read_delimited(data: bytes, delimiter: str) -> list[RawRow]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for row_number, row in enumerate(reader, start=2):
        values = {column: cell_to_text(value) for column, value in row.items() if column is not None}
        rows.append(RawRow(locator=Locator(row=row_number), values=values))
    return rows


def read_csv(data: bytes) -> list[RawRow]:
    return read_delimited(data, ",")


def read_tsv(data: bytes) -> list[RawRow]:
    return read_delimited(data, "\t")


def objects_to_rows(objects: list[dict]) -> list[RawRow]:
    rows = []
    for row_number, item in enumerate(objects, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"row {row_number} is not an object")
        values = {str(key): cell_to_text(value) for key, value in item.items()}
        rows.append(RawRow(locator=Locator(row=row_number), values=values))
    return rows


def read_json(data: bytes) -> list[RawRow]:
    objects = json.loads(data.decode("utf-8-sig"))
    if not isinstance(objects, list):
        raise ValueError("a JSON export must be an array of objects")
    return objects_to_rows(objects)


def read_jsonl(data: bytes) -> list[RawRow]:
    lines = data.decode("utf-8-sig").splitlines()
    objects = [json.loads(line) for line in lines if line.strip()]
    return objects_to_rows(objects)


def read_xlsx(data: bytes) -> list[RawRow]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    sheet_rows = sheet.iter_rows(values_only=True)
    header = [cell_to_text(cell) for cell in next(sheet_rows, [])]
    rows = []
    for row_number, cells in enumerate(sheet_rows, start=2):
        values = {column: cell_to_text(cell) for column, cell in zip(header, cells) if column}
        if not any(values.values()):
            continue
        locator = Locator(sheet=sheet.title, row=row_number)
        rows.append(RawRow(locator=locator, values=values))
    workbook.close()
    return rows
